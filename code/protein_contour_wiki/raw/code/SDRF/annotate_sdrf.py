"""
Phase 4 v3: 注释生成 SDRF xlsx,GelBand 行黄色高亮
Output: out/PXD023690_SDRF_annotated.xlsx
"""
import json
import re
import openpyxl
from openpyxl.styles import PatternFill
from pathlib import Path

OUT_DIR = Path("./out")
TEMPLATE = Path("SDRF.xlsx")

files = json.loads((OUT_DIR / "raw_files.json").read_text())
meta = json.loads((OUT_DIR / "project_metadata.json").read_text())

# ============ 字典 / 常量 ============

ORGANISM_MAP = {"mouse": "mus musculus (mouse)", "human": "homo sapiens (human)"}

MOD_FIXED = "fixed=unimod 4"
MOD_VAR_TOTAL = "variable=unimod 35, unimod 1"
MOD_VAR_PHOSPHO = "variable=unimod 35, unimod 1, unimod 21"
LABEL_TOTAL = "Carbamidomethyl C; Oxidation M; Acetyl N-term"
LABEL_PHOSPHO = "Carbamidomethyl C; Oxidation M; Acetyl N-term; Phospho STY"
TPP_TOTAL, TPP_PHOSPHO = "TP", "P"
FRAC2COMP = {1: "Cyt", 2: "Cyt", 3: "Mem", 4: "Mem", 5: "Nuc", 6: "Nuc"}

CELL_PATTERNS = ["HeLa", "U2OS"]
ORGAN_PART_PATTERNS = ["Kidney", "Liver", "Muscle", "Brain", "Heart", "Lung", "Spleen"]

STIM_RE = re.compile(r"^(EGF|OSMSTRS{0,2}|OSM|TGF|IL\d*)$", re.I)
TIME_RE = re.compile(r"^(\d+(min|h|hour)|CTRL|control)$", re.I)
COMBINED_TIME_RE = re.compile(r"^(EGF|OSMSTRS{0,2}|OSM|TGF)(\d+(min|h|hour)|CTRL)$", re.I)
OSMSTR_TIME_RE = re.compile(r"^\d+h(SORB|REL|RELEASE).*$", re.I)
RELEASE_TOKEN_RE = re.compile(r"^\d+min(REL|RELEASE)$", re.I)
FRAC_RE = re.compile(r"^(?:fraction|frac|FR|F)(\d+)$", re.I)
REP_RE = re.compile(r"^Rep\d+$", re.I)
TRAIL_DIGIT_RE = re.compile(r"^\d{2,3}$")
WT_ID_RE = re.compile(r"^WT-(Stim-)?\d+[LR]?$")
GELBAND_COORD_RE = re.compile(r"^[A-J]\d{1,2}$")
GELBAND_MARKER_RE = re.compile(r"^GelBand$", re.I)
PERC_RE = re.compile(r"^(\d+)perc$", re.I)
TREATMENT_LITERAL = {"stimulated", "serumstarv"}

GELBAND_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Excel warning yellow


def parse_time(toks):
    for tok in toks:
        if COMBINED_TIME_RE.match(tok):
            return tok
    for i, tok in enumerate(toks):
        if STIM_RE.match(tok) and i + 1 < len(toks):
            nxt = toks[i + 1]
            if TIME_RE.match(nxt) or OSMSTR_TIME_RE.match(nxt) or RELEASE_TOKEN_RE.match(nxt):
                return f"{tok}_{nxt}"
    for tok in toks:
        if tok.lower() in ("ctrl", "control"):
            return "control"
    for tok in toks:
        if tok.lower() in TREATMENT_LITERAL:
            return tok.lower().replace("serumstarv", "serum_starved")
    # (Patch 8) WT-Stim-* / WT-* (mouse muscle individual ID)
    for tok in toks:
        if re.match(r"^WT-Stim-\d+[LR]?$", tok):
            return "stimulated"
        if re.match(r"^WT-\d+[LR]?$", tok):
            return "control"
    for tok in toks:
        if STIM_RE.match(tok):
            return f"{tok}_treated"
    return ""


def parse_filename(fname):
    name = fname.rsplit(".raw", 1)[0]
    toks = name.split("_")
    toks_lower = [t.lower() for t in toks]
    info = {"organism": "human", "organism_part": "", "cell": "",
            "modtype": "total", "exp_design": "", "time": "", "compartment": "",
            "bio_rep": "", "tech_rep": "", "wt_id": "", "is_gelband": False}

    # GelBand 检测优先(影响后续逻辑)
    if any(GELBAND_MARKER_RE.match(t) for t in toks):
        info["is_gelband"] = True
        info["exp_design"] = "GelBand"
        # 提取百分比和坐标
        perc = next((m.group(1) for t in toks if (m := PERC_RE.match(t))), "")
        coord = next((t for t in toks if GELBAND_COORD_RE.match(t)), "")
        comp_parts = ["GelBand"]
        if perc: comp_parts.append(f"{perc}perc")
        if coord: comp_parts.append(coord)
        info["compartment"] = "_".join(comp_parts)
        # GelBand 文件:modtype 默认 total(等博后确认),其他字段尽力而为
        # 不 return,继续解析其他字段(可能漏标的 cell/organism)

    # Organism
    if any(t in ("mouse", "mus") for t in toks_lower):
        info["organism"] = "mouse"

    # Organism part
    if info["organism"] == "mouse":
        for op in ORGAN_PART_PATTERNS:
            if any(op.lower() == t for t in toks_lower):
                info["organism_part"] = op
                break

    # Cell type
    if info["organism"] == "human":
        for cp in CELL_PATTERNS:
            if any(cp.lower() == t for t in toks_lower):
                info["cell"] = cp
                break

    # Modification type
    if any("phos" in t for t in toks_lower):
        info["modtype"] = "phospho"
    elif any(t in ("prot", "wcl", "total") for t in toks_lower):
        info["modtype"] = "total"
    # GelBand 文件如果没标 phos/prot,保持 default total

    # Experimental design (非 GelBand)
    if not info["is_gelband"]:
        for t in toks:
            if t.lower() == "subcell":
                info["exp_design"] = "SUBCELL"; break
            if t.lower() == "wcl":
                info["exp_design"] = "WCL"; break

    # Time
    info["time"] = parse_time(toks)

    # WT individual ID (mouse muscle)
    info["wt_id"] = next((t for t in toks if WT_ID_RE.match(t)), "")

    # Fraction → compartment(GelBand 已经处理过,跳过)
    if not info["is_gelband"]:
        if info["exp_design"] == "WCL":
            info["compartment"] = "WholeCell"
        else:
            for tok in toks:
                m = FRAC_RE.match(tok)
                if m:
                    n = int(m.group(1))
                    info["compartment"] = FRAC2COMP.get(n, "")
                    break

    # Replicate 双标
    explicit = next((t for t in toks if REP_RE.match(t)), None)
    trailing = toks[-1] if toks and TRAIL_DIGIT_RE.match(toks[-1]) else None
    if explicit and trailing:
        info["bio_rep"] = explicit.lower()
        info["tech_rep"] = trailing
    elif explicit:
        info["bio_rep"] = explicit.lower()
    elif trailing:
        info["bio_rep"] = f"rep{int(trailing)}"

    return info


def make_source_name(info):
    parts = []
    if info["is_gelband"]:
        parts.append(info["compartment"])  # GelBand_25perc_C3
    elif info["cell"]:
        parts.append(info["cell"])
    elif info["organism_part"]:
        parts.append(f"Mouse{info['organism_part']}")
    if info["wt_id"]:
        parts.append(info["wt_id"])
    if info["time"]:
        parts.append(info["time"])
    parts.append("Phos" if info["modtype"] == "phospho" else "Prot")
    if info["bio_rep"]:
        parts.append(info["bio_rep"])
    return "_".join(parts) if parts else "sample"


# ============ 写 xlsx ============

wb = openpyxl.load_workbook(TEMPLATE)
ws = wb["Sheet1"]

# 找两个 mod param 列
mod_param_cols = [c.column for c in ws[1] if c.value == "comment[modification parameter]"]
assert len(mod_param_cols) == 2

# 修正 typo
for cell in ws[1]:
    if cell.value == "comment[proteomics data adquisition method]":
        cell.value = "comment[proteomics data acquisition method]"

# 删除示例行
ws.delete_rows(2, ws.max_row - 1)

# 列名 → 列号(去重)
header = {}
for cell in ws[1]:
    if cell.value and cell.value != "comment[modification parameter]":
        header[cell.value] = cell.column

# 新增 2 列
last_col = ws.max_column
ws.cell(row=1, column=last_col + 1, value="comment[modification labels]")
ws.cell(row=1, column=last_col + 2, value="comment[TP/P]")
header["comment[modification labels]"] = last_col + 1
header["comment[TP/P]"] = last_col + 2

# Instrument
instr = ""
for i in meta.get("instruments", []):
    if isinstance(i, dict):
        instr = i.get("name") or i.get("accession") or str(i)
    else:
        instr = str(i)
    break

# 写入 + 统计
gelband_count = 0
mouse_count = 0
human_count = 0
phos_count = 0
total_count = 0

for i, f in enumerate(files, start=2):
    info = parse_filename(f["fileName"])

    # 统计
    if info["organism"] == "mouse":
        mouse_count += 1
    else:
        human_count += 1
    if info["modtype"] == "phospho":
        phos_count += 1
    else:
        total_count += 1

    cells = {
        "source name": make_source_name(info),
        "characteristics[organism]": ORGANISM_MAP[info["organism"]],
        "characteristics[organism part]": info["organism_part"],
        "characteristics[cell type]": info["cell"],
        "characteristics[sampling time]": info["time"],
        "characteristics[biological replicate]": info["bio_rep"],
        "comment[instrument]": instr,
        "comment[fraction identifier]": info["compartment"],
        "comment[proteomics data acquisition method]": "DIA",
        "comment[cleavage agent details]": "Trypsin",
        "comment[file uri]": f["uri"],
        "comment[technical replicate]": info["tech_rep"],
        "comment[data file]": f["fileName"],
        "comment[modification labels]": LABEL_PHOSPHO if info["modtype"] == "phospho" else LABEL_TOTAL,
        "comment[TP/P]": TPP_PHOSPHO if info["modtype"] == "phospho" else TPP_TOTAL,
    }
    for col_name, val in cells.items():
        if col_name in header:
            ws.cell(row=i, column=header[col_name], value=val)

    ws.cell(row=i, column=mod_param_cols[0], value=MOD_FIXED)
    var_str = MOD_VAR_PHOSPHO if info["modtype"] == "phospho" else MOD_VAR_TOTAL
    ws.cell(row=i, column=mod_param_cols[1], value=var_str)

    # GelBand 高亮
    if info["is_gelband"]:
        gelband_count += 1
        for col in range(1, ws.max_column + 1):
            ws.cell(row=i, column=col).fill = GELBAND_FILL

OUT = OUT_DIR / "PXD023690_SDRF_annotated.xlsx"
wb.save(OUT)

print(f"✅ Phase 4 v3 done.")
print(f"   Output: {OUT}")
print()
print("=" * 50)
print("STATS")
print("=" * 50)
print(f"  Total rows:       {len(files)}")
print(f"  Mouse:            {mouse_count}")
print(f"  Human:            {human_count}")
print(f"  Phospho:          {phos_count}")
print(f"  Total proteome:   {total_count}")
print(f"  GelBand (yellow): {gelband_count}")
print("=" * 50)
